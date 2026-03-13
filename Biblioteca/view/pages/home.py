import flet as ft
from model.database import Database

import os
import sys
import tempfile
import re
import subprocess
from datetime import datetime, timedelta
from openpyxl import Workbook
import calendar
import math


class HomePage(ft.Column):
    def __init__(self, navigate, page: ft.Page, query_params: dict = None):
        super().__init__()
        self.expand = True
        self.scroll = ft.ScrollMode.AUTO
        self._page = page
        self.navigate = navigate
        self._db = Database()
        self._estado_actual = "TODOS"
        self._page_size = 5
        
        # Restaurar página desde query params si existen
        initial_page = 1
        if query_params and isinstance(query_params, dict) and "page" in query_params:
            try:
                initial_page = max(1, int(query_params["page"]))
            except (ValueError, TypeError):
                initial_page = 1
        
        self._pagina_actual = initial_page

        # =========================
        # Helpers
        # =========================
        def snack(msg: str):
            self._page.snack_bar = ft.SnackBar(ft.Text(msg))
            self._page.snack_bar.open = True
            self._page.update()

        def safe_filename(name: str) -> str:
            name = (name or "").strip()
            name = re.sub(r"[^\w\-. ]+", "", name, flags=re.UNICODE)
            if not name.lower().endswith(".pdf"):
                name += ".pdf"
            return name[:150] if name else "documento.pdf"

        def open_file_default_app(path: str):
            try:
                if sys.platform.startswith("win"):
                    os.startfile(path)
                elif sys.platform == "darwin":
                    os.system(f'open "{path}"')
                else:
                    os.system(f'xdg-open "{path}"')
            except Exception as ex:
                print("[ERROR] open_file_default_app:", ex)
                snack(f"No se pudo abrir el archivo: {ex}")

        # =========================
        # Ver PDF
        # =========================
        def ver_pdf(id_libro: int):
            try:
                data = self._db.get_libro_pdf(int(id_libro))
                if not data:
                    snack("Este libro no tiene PDF aún.")
                    return

                nombre = (data.get("nombre_archivo") or f"libro_{id_libro}.pdf")
                contenido = data.get("contenido")
                if contenido is None:
                    snack("El PDF existe pero vino vacío desde la BD.")
                    return

                pdf_bytes = bytes(contenido)

                out_path = os.path.join(
                    tempfile.gettempdir(),
                    f"libro_{id_libro}_{safe_filename(nombre)}"
                )

                with open(out_path, "wb") as f:
                    f.write(pdf_bytes)

                open_file_default_app(out_path)

            except Exception as ex:
                print("[ERROR] ver_pdf:", ex)
                snack(f"Error abriendo PDF: {ex}")

        # =========================
        # Picker PDF
        # =========================
        def pick_pdf_windows() -> str | None:
            """Abre el selector de archivos rápidamente. Primero intenta Tkinter (rápido),
            si falla, usa PowerShell como respaldo."""
            # 1) Intento con Tkinter (rápido y directo)
            try:
                from tkinter import Tk
                from tkinter.filedialog import askopenfilename

                root = Tk()
                root.withdraw()
                root.attributes("-topmost", True)
                initial_dir = os.path.join(os.path.expanduser("~"), "Documents")
                if not os.path.isdir(initial_dir):
                    initial_dir = os.path.expanduser("~")
                path = askopenfilename(
                    title="Selecciona un archivo PDF",
                    filetypes=[("PDF files", "*.pdf")],
                    initialdir=initial_dir,
                )
                root.destroy()
                if path:
                    return path
            except Exception:
                # Continuar a PowerShell
                pass

            # 2) Respaldo con PowerShell (sin archivo temporal, comando inline)
            try:
                powershell_51 = r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"
                creationflags = 0
                if sys.platform.startswith("win"):
                    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

                exe = powershell_51 if os.path.exists(powershell_51) else "powershell.exe"
                ps_command = (
                    "Add-Type -AssemblyName System.Windows.Forms;"
                    "$ofd = New-Object System.Windows.Forms.OpenFileDialog;"
                    "$ofd.Filter = \"Archivos PDF (.pdf)|*.pdf\";"
                    "$ofd.Multiselect = $false;"
                    "$ofd.Title = \"Selecciona un archivo PDF\";"
                    "$null = $ofd.ShowDialog();"
                    "if ($ofd.FileName) { Write-Output $ofd.FileName }"
                )

                result = subprocess.run(
                    [exe, "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-Command", ps_command],
                    capture_output=True,
                    text=True,
                    timeout=30,
                    creationflags=creationflags,
                )

                if result.returncode == 0:
                    path = (result.stdout or "").strip()
                    if path:
                        return path

                err = (result.stderr or "").strip()
                if err:
                    snack(f"PowerShell no pudo abrir selector: {err[:180]}")

            except Exception as ex:
                snack(f"PowerShell falló: {ex}")

            return None

        # =========================
        # UI
        # =========================
        button_crear_libro = ft.ElevatedButton(
            content=ft.Row([
                ft.Icon(ft.Icons.ADD_ROUNDED, size=20),
                ft.Text("Crear libro", size=14, weight=ft.FontWeight.W_500)
            ], spacing=8),
            on_click=self.abrir_dialogo_crear_libro,
            bgcolor="#1976d2",
            color=ft.Colors.WHITE,
            height=48,
            style=ft.ButtonStyle(
                shape=ft.RoundedRectangleBorder(radius=10),
                elevation=2,
            ),
        )

        search_input = ft.TextField(
            hint_text="Buscar por libro, código o autor...",
            prefix_icon=ft.Icons.SEARCH_ROUNDED,
            width=500,
            height=50,
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            border_color="#e0e0e0",
            focused_border_color="#1976d2",
            focused_border_width=2,
            text_size=14,
            content_padding=ft.padding.only(left=15, right=15, top=10, bottom=10),
        )

        # Filtro de estado (Dropdown en lugar de botones)
        def set_estado(estado: str):
            self._estado_actual = estado
            aplicar_filtros(q=search_input.value, estado=self._estado_actual, reset_pagina=True)

        estado_dd = ft.Dropdown(
            width=200,
            height=50,
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            border_color="#e0e0e0",
            focused_border_color="#1976d2",
            value="TODOS",
            label="Filtrar",
            text_size=14,
            options=[
                ft.dropdown.Option("TODOS", "Todos"),
                ft.dropdown.Option("ACTIVOS", "Activos"),
                ft.dropdown.Option("INACTIVOS", "Inactivos"),
            ],
        )
        # asignar handler luego (compatibilidad Flet)
        estado_dd.on_change = lambda e: set_estado(e.control.value)

        # Filtro de tipo (Digital / Físico)
        def set_tipo(tipo: str):
            self._tipo_actual = tipo
            aplicar_filtros(q=search_input.value, estado=self._estado_actual, tipo=tipo, reset_pagina=True)

        tipo_dd = ft.Dropdown(
            width=140,
            height=50,
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            border_color="#e0e0e0",
            focused_border_color="#1976d2",
            value="TODOS",
            label="Tipo",
            text_size=14,
            options=[
                ft.dropdown.Option("TODOS", "Todos"),
                ft.dropdown.Option("DIGITAL", "Digital"),
                ft.dropdown.Option("FISICO", "Físico"),
            ],
        )
        tipo_dd.on_change = lambda e: set_tipo(e.control.value)

        libros_table = ft.DataTable(
            bgcolor=ft.Colors.WHITE,
            border=ft.border.all(1, "#e0e0e0"),
            border_radius=12,
            heading_row_color="#f5f7fa",
            heading_row_height=56,
            data_row_min_height=60,
            data_row_max_height=65,
            column_spacing=48,
            horizontal_margin=20,
            divider_thickness=0.5,
            columns=[
                ft.DataColumn(label=ft.Text("Código de barras", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),

                ft.DataColumn(label=ft.Text("Título", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("Clasificación DUI", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("Autores", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("ISBN", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("Tipo", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("Documentos", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
                ft.DataColumn(label=ft.Text("Acciones", text_align=ft.TextAlign.CENTER, weight=ft.FontWeight.BOLD, color="#0d47a1")),
            ],
            rows=[],
        )

        self._libros_cache = []
        self._libros_filtrados = []

        # =========================
        # Subir PDF
        # =========================
        def subir_pdf_para_libro(id_libro: int):
            path = pick_pdf_windows()
            if not path:
                return

            if not path.lower().endswith(".pdf"):
                snack("Solo se permiten archivos PDF.")
                return

            with open(path, "rb") as f:
                pdf_bytes = f.read()

            self._db.upsert_libro_pdf(
                id_libro=id_libro,
                nombre_archivo=safe_filename(os.path.basename(path)),
                contenido=pdf_bytes,
            )

            snack("PDF guardado correctamente ✅")
        # =========================
        # Mostrar libros
        # =========================
        def mostrar_libros(e=None):
            libros_table.rows.clear()
            self._libros_cache = self._db.get_libros()
            try:
                actualizar_stats()
            except Exception:
                pass
            aplicar_filtros(reset_pagina=True)

        def mostrar_libros_sin_reset(e=None):
            """Recarga los libros manteniendo la página actual (para después de editar/crear)"""
            libros_table.rows.clear()
            self._libros_cache = self._db.get_libros()
            try:
                actualizar_stats()
            except Exception:
                pass
            aplicar_filtros(reset_pagina=False)

        self._mostrar_libros = mostrar_libros
        self._mostrar_libros_sin_reset = mostrar_libros_sin_reset

        # =========================
        # Helper para navegar manteniendo página
        # =========================
        def navigate_with_page(route: str):
            """Navega a una ruta manteniendo el query param de página actual"""
            if self._pagina_actual > 1:
                route = f"{route}?page={self._pagina_actual}"
            navigate(route)
        
        self._navigate_with_page = navigate_with_page

        # =========================
        # Toggle activo/inactivo
        # =========================
        def toggle_activo(libro_id: int, estado_actual: bool):
            nuevo_estado = 0 if estado_actual else 1
            try:
                self._db.set_libro_activo(libro_id, nuevo_estado)
                snack("Libro activado ✅" if nuevo_estado == 1 else "Libro desactivado ✅")
                mostrar_libros()
            except Exception as ex:
                snack(f"Error al actualizar estado: {ex}")

        # =========================
        # Construir filas
        # =========================
        def build_row(libro: dict) -> ft.DataRow:
            id_libro = int(libro["id_libro"])
            tipo = (libro.get("tipo") or "").strip().upper()

            # ---- Documentos (botones compactos con color)
            if tipo == "DIGITAL":
                btn_cargar = ft.Container(
                    content=ft.Row([ft.Icon(ft.Icons.UPLOAD_FILE, color=ft.Colors.WHITE), ft.Text("Cargar", size=12, color=ft.Colors.WHITE)], spacing=6),
                    on_click=lambda e, i=id_libro: subir_pdf_para_libro(i),
                    bgcolor="#1976d2",
                    padding=ft.padding.symmetric(horizontal=10),
                    height=36,
                    border_radius=8,
                    alignment=ft.Alignment.CENTER,
                )

                btn_ver = ft.Container(
                    content=ft.Row([ft.Icon(ft.Icons.PICTURE_AS_PDF, color=ft.Colors.WHITE), ft.Text("Ver", size=12, color=ft.Colors.WHITE)], spacing=6),
                    on_click=lambda e, i=id_libro: ver_pdf(i),
                    bgcolor="#00838f",
                    padding=ft.padding.symmetric(horizontal=10),
                    height=36,
                    border_radius=8,
                    alignment=ft.Alignment.CENTER,
                )

                documentos = ft.DataCell(
                    ft.Row(
                        [btn_cargar, btn_ver],
                        spacing=8,
                        alignment=ft.MainAxisAlignment.CENTER,
                        expand=True,
                    )
                )
            else:
                documentos = ft.DataCell(
                    ft.Row(
                        [ft.Text("No aplica", color=ft.Colors.GREY)],
                        alignment=ft.MainAxisAlignment.CENTER,
                        expand=True,
                    )
                )

            def action_button(icon, bgcolor, tooltip, on_click=None):
                return ft.Container(
                    width=36,
                    height=36,
                    bgcolor=bgcolor,
                    border_radius=6,
                    alignment=ft.Alignment.CENTER,
                    tooltip=tooltip,
                    on_click=on_click,
                    content=ft.Icon(icon, color=ft.Colors.WHITE, size=18),
                )

            # ---- Activo robusto
            activo_value = libro.get("activo")
            is_activo = (
                activo_value == 1
                or activo_value is True
                or str(activo_value).strip().lower() in ("1", "true")
            )

            # ---- Acciones
            acciones = ft.DataCell(
                ft.Row(
                    [
                        action_button(
                            ft.Icons.VISIBILITY,
                            ft.Colors.BLUE,
                            "Detalles",
                            lambda e, i=id_libro: self._navigate_with_page(f"/libro/{i}"),
                        ),
                        action_button(
                            ft.Icons.EDIT,
                            ft.Colors.ORANGE,
                            "Modificar",
                            lambda e, i=id_libro: self.abrir_dialogo_editar_libro(i),
                        ),
                        action_button(
                            ft.Icons.CHECK_CIRCLE if is_activo else ft.Icons.BLOCK,
                            ft.Colors.GREEN if is_activo else ft.Colors.RED,
                            "Desactivar" if is_activo else "Activar",
                            lambda e, i=id_libro, st=is_activo: toggle_activo(i, st),
                        ),
                    ],
                    spacing=10,
                    alignment=ft.MainAxisAlignment.CENTER,
                    expand=True,
                )
            )

            autores = (libro.get("autores", "") or "").strip()
            autores_short = (autores[:60] + "…") if len(autores) > 60 else autores

            codigo_barras = (libro.get("codigo_barras", "") or "").strip()
            codigo_barras_short = (codigo_barras[:40] + "…") if len(codigo_barras) > 40 else codigo_barras

            # Celda de código de barras: texto normal (sin badge)
            codigo_cell = ft.DataCell(
                ft.Container(
                    content=ft.Text(codigo_barras_short, text_align=ft.TextAlign.LEFT, size=12),
                    alignment=ft.Alignment(-1, 0),
                    height=60,
                )
            )

            # Celda tipo con estilo idéntico a 'estado' en reservas (bg sólido, texto blanco)
            if tipo == "DIGITAL":
                # color distinto para Digital (resaltado solicitado)
                tipo_bg = "#ff7043"
            elif tipo == "FISICO":
                tipo_bg = "#6a1b9a"
            else:
                tipo_bg = "#9e9e9e"

            tipo_cell = ft.DataCell(
                ft.Container(
                    content=ft.Row([
                        ft.Text(tipo.title(), color=ft.Colors.WHITE, size=12, weight=ft.FontWeight.BOLD),
                    ], tight=True),
                    bgcolor=tipo_bg,
                    padding=ft.padding.symmetric(horizontal=12, vertical=4),
                    height=36,
                    border_radius=8,
                    alignment=ft.Alignment.CENTER,
                )
            )

            # Categoría (para mostrar como subtítulo bajo el título)
            categoria = (libro.get("categoria", "") or "").strip()
            categoria_short = (categoria[:60] + "…") if len(categoria) > 60 else categoria

            # Título + categoría en una sola celda (similar a reservas: nombre + cédula)
            titulo = (libro.get("titulo", "") or "").strip()
            title_cell = ft.DataCell(
                ft.Container(
                    content=ft.Column([
                        ft.Text(titulo, text_align=ft.TextAlign.LEFT),
                        ft.Text(categoria_short, size=12, color="#757575", text_align=ft.TextAlign.LEFT),
                    ], spacing=4, alignment=ft.MainAxisAlignment.START, horizontal_alignment=ft.CrossAxisAlignment.START),
                    alignment=ft.Alignment(-1, 0),
                    height=60,
                    padding=ft.padding.only(top=6, bottom=6),
                    width=150,
                )
            )

            return ft.DataRow(
                cells=[
                    codigo_cell,
                    title_cell,
                    ft.DataCell(ft.Container(content=ft.Text(libro.get("clasificacion_dui", "") or "", text_align=ft.TextAlign.LEFT), alignment=ft.Alignment(-1, 0), height=60)),
                    ft.DataCell(ft.Container(content=ft.Text(autores_short, text_align=ft.TextAlign.LEFT), alignment=ft.Alignment(-1, 0), height=60)),
                    ft.DataCell(ft.Container(content=ft.Text(libro.get("isbn", ""), text_align=ft.TextAlign.LEFT), alignment=ft.Alignment(-1, 0), height=60)),
                    tipo_cell,
                    documentos,
                    acciones,
                ]
            )

        # =========================
        # Filtros combinados
        # =========================
        def libro_match_estado(libro: dict, estado: str) -> bool:
            activo_value = libro.get("activo")
            is_activo = (
                activo_value == 1
                or activo_value is True
                or str(activo_value).strip().lower() in ("1", "true")
            )

            estado = (estado or "TODOS").strip().upper()
            if estado == "ACTIVOS":
                return is_activo
            if estado == "INACTIVOS":
                return not is_activo
            return True

        def libro_match_tipo(libro: dict, tipo: str) -> bool:
            t = ((libro.get("tipo") or "").strip() or "").upper()
            tipo = (tipo or "TODOS").strip().upper()
            if tipo == "TODOS":
                return True
            if tipo == "DIGITAL":
                return t == "DIGITAL"
            if tipo == "FISICO":
                return t == "FISICO"
            return True

        def libro_match_texto(libro: dict, q: str) -> bool:
            if not q:
                return True

            activo_value = libro.get("activo")
            is_activo = (
                activo_value == 1
                or activo_value is True
                or str(activo_value).strip().lower() in ("1", "true")
            )

            valores = [
                libro.get("codigo_barras", ""),  # ✅ NUEVO (busqueda)
                libro.get("titulo", ""),
                libro.get("clasificacion_dui", ""),
                libro.get("categoria", ""),
                libro.get("autores", ""),
                libro.get("descripcion", ""),
                libro.get("isbn", ""),
                libro.get("tipo", ""),
                "si" if is_activo else "no",
                "activo" if is_activo else "inactivo",
            ]

            haystack = " | ".join([(v or "") for v in valores]).lower()
            return q in haystack

        # =========================
        # Paginación
        # =========================
        pagination_label = ft.Text("Página 1 de 1", size=12, color="#546e7a")

        def change_page(delta: int):
            total = max(1, math.ceil(len(self._libros_filtrados) / self._page_size))
            self._pagina_actual = min(max(1, self._pagina_actual + delta), total)
            aplicar_filtros()

        btn_prev = ft.IconButton(
            icon=ft.Icons.CHEVRON_LEFT,
            icon_color="#546e7a",
            tooltip="Anterior",
            on_click=lambda e: change_page(-1),
        )
        btn_next = ft.IconButton(
            icon=ft.Icons.CHEVRON_RIGHT,
            icon_color="#546e7a",
            tooltip="Siguiente",
            on_click=lambda e: change_page(1),
        )

        # ✅ aceptar valores directos (evita leer value viejo)
        def aplicar_filtros(q=None, estado=None, tipo=None, reset_pagina: bool = False):
            libros_table.rows.clear()

            q = (q if q is not None else (search_input.value or "")).strip().lower()
            estado = (estado if estado is not None else (getattr(estado_dd, 'value', None) or self._estado_actual or "TODOS")).strip().upper()
            tipo = (tipo if tipo is not None else (getattr(tipo_dd, 'value', None) or getattr(self, '_tipo_actual', 'TODOS') or "TODOS")).strip().upper()

            filtrados = []
            for libro in self._libros_cache:
                if libro_match_estado(libro, estado) and libro_match_texto(libro, q) and libro_match_tipo(libro, tipo):
                    filtrados.append(libro)

            self._libros_filtrados = filtrados

            total_pages = max(1, math.ceil(len(filtrados) / self._page_size))
            if reset_pagina:
                self._pagina_actual = 1
            if self._pagina_actual > total_pages:
                self._pagina_actual = total_pages

            start = (self._pagina_actual - 1) * self._page_size
            end = start + self._page_size
            pagina_items = filtrados[start:end]

            for libro in pagina_items:
                libros_table.rows.append(build_row(libro))

            pagination_label.value = f"Página {self._pagina_actual} de {total_pages}"
            btn_prev.disabled = self._pagina_actual <= 1
            btn_next.disabled = self._pagina_actual >= total_pages

            self._page.update()

        # ✅ usar e.control.value
        def on_search_change(e):
            if not self._libros_cache:
                self._libros_cache = self._db.get_libros()
            aplicar_filtros(q=e.control.value, estado=self._estado_actual, reset_pagina=True)

        search_input.on_change = on_search_change

        # =========================
        # Exportar Excel (Semana/Mes)
        # =========================
        def parse_fecha(valor):
            if not valor:
                return None
            if isinstance(valor, datetime):
                return valor
            s = str(valor).strip()
            formatos = [
                "%Y-%m-%d",
                "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y",
                "%d/%m/%Y %H:%M:%S",
            ]
            for fmt in formatos:
                try:
                    return datetime.strptime(s, fmt)
                except Exception:
                    pass
            try:
                return datetime.fromisoformat(s)
            except Exception:
                return None

        def get_fecha_registro(libro: dict):
            for key in [
                "fecha_registro",
                "fecha_creacion",
                "fecha",
                "fecha_alta",
                "creado_en",
            ]:
                if key in libro and libro[key] is not None:
                    return parse_fecha(libro[key])
            return None

        def filtrar_por_periodo(libros, periodo: str):
            ahora = datetime.now()
            if periodo == "SEMANA":
                inicio = ahora - timedelta(days=7)
            else:  # MES
                inicio = ahora - timedelta(days=30)

            filtrados = []
            tiene_fecha = False
            for l in libros:
                f = get_fecha_registro(l)
                if f:
                    tiene_fecha = True
                    if f >= inicio:
                        filtrados.append(l)
            # Si no hay fechas en ningún registro, devolver todos para ese estado
            return filtrados if tiene_fecha else libros

        def exportar_excel_estado(periodo: str, estado: str):
            if not self._libros_cache:
                self._libros_cache = self._db.get_libros()

            estado = estado.strip().upper()
            periodo = periodo.strip().upper()

            seleccion = [l for l in self._libros_cache if libro_match_estado(l, estado)]
            seleccion = filtrar_por_periodo(seleccion, periodo)

            # Preparar Excel
            wb = Workbook()
            ws = wb.active
            ws.title = f"{estado.title()}-{periodo.title()}"

            headers = [
                "Código de barras",
                "Título",
                "Clasificación DUI",
                "categoria",
                "Autores",
                "ISBN",
                "Tipo",
                "Activo",
            ]
            ws.append(headers)

            # preparar mapa de categorias
            categorias_map = {}
            try:
                cats = self._db.get_categorias()
                for c in cats:
                    cid = c.get("id_categoria") or c.get("id")
                    if cid is not None:
                        categorias_map[str(cid)] = c.get("nombre") or c.get("nombre_categoria") or ""
            except Exception:
                categorias_map = {}

            for libro in seleccion:
                activo_value = libro.get("activo")
                is_activo = (
                    activo_value == 1
                    or activo_value is True
                    or str(activo_value).strip().lower() in ("1", "true")
                )
                # intentar resolver nombre de categoria si viene id
                cat_val = libro.get("id_categoria") if libro.get("id_categoria") is not None else libro.get("categoria")
                cat_name = categorias_map.get(str(cat_val), libro.get("categoria") or "")
                ws.append([
                    (libro.get("codigo_barras") or ""),
                    (libro.get("titulo") or ""),
                    (libro.get("clasificacion_dui") or ""),
                    (cat_name),
                    (libro.get("autores") or ""),
                    (libro.get("isbn") or ""),
                    (libro.get("tipo") or ""),
                    ("ACTIVO" if is_activo else "INACTIVO"),
                ])

            nombre = f"libros_{estado.lower()}_{periodo.lower()}.xlsx"
            out_path = os.path.join(tempfile.gettempdir(), nombre)
            try:
                wb.save(out_path)
                snack(f"Excel generado: {nombre}")
                open_file_default_app(out_path)
            except Exception as ex:
                snack(f"No se pudo generar Excel: {ex}")

        # ----- Exportes por rango explícito -----
        def rango_semana_actual():
            hoy = datetime.now()
            delta = hoy.weekday()  # lunes=0
            inicio = datetime(hoy.year, hoy.month, hoy.day) - timedelta(days=delta)
            fin = inicio + timedelta(days=6, hours=23, minutes=59, seconds=59)
            return inicio, fin

        def rango_semana_anterior():
            inicio_actual, _ = rango_semana_actual()
            inicio = inicio_actual - timedelta(days=7)
            fin = inicio + timedelta(days=6, hours=23, minutes=59, seconds=59)
            return inicio, fin

        def rango_mes_actual():
            hoy = datetime.now()
            inicio = datetime(hoy.year, hoy.month, 1)
            dias = calendar.monthrange(hoy.year, hoy.month)[1]
            fin = datetime(hoy.year, hoy.month, dias, 23, 59, 59)
            return inicio, fin

        def rango_mes_anterior():
            hoy = datetime.now()
            year = hoy.year
            month = hoy.month - 1
            if month == 0:
                month = 12
                year -= 1
            inicio = datetime(year, month, 1)
            dias = calendar.monthrange(year, month)[1]
            fin = datetime(year, month, dias, 23, 59, 59)
            return inicio, fin

        def rango_mes_especifico(year: int, month: int):
            inicio = datetime(year, month, 1)
            dias = calendar.monthrange(year, month)[1]
            fin = datetime(year, month, dias, 23, 59, 59)
            return inicio, fin

        def filtrar_por_rango(libros, inicio: datetime, fin: datetime):
            filtrados = []
            tiene_fecha = False
            for l in libros:
                f = get_fecha_registro(l)
                if f:
                    tiene_fecha = True
                    if inicio <= f <= fin:
                        filtrados.append(l)
            return filtrados if tiene_fecha else libros

        def exportar_excel_estado_rango(estado: str, inicio: datetime, fin: datetime, nombre_sufijo: str):
            if not self._libros_cache:
                self._libros_cache = self._db.get_libros()

            estado = estado.strip().upper()
            seleccion = [l for l in self._libros_cache if libro_match_estado(l, estado)]
            seleccion = filtrar_por_rango(seleccion, inicio, fin)

            wb = Workbook()
            ws = wb.active
            ws.title = f"{estado.title()}"

            headers = [
                "Código de barras",
                "Título",
                "Clasificación DUI",
                "categoria",
                "Autores",
                "ISBN",
                "Tipo",
                "Activo",
            ]
            ws.append(headers)

            # preparar mapa de categorias
            categorias_map = {}
            try:
                cats = self._db.get_categorias()
                for c in cats:
                    cid = c.get("id_categoria") or c.get("id")
                    if cid is not None:
                        categorias_map[str(cid)] = c.get("nombre") or c.get("nombre_categoria") or ""
            except Exception:
                categorias_map = {}

            for libro in seleccion:
                activo_value = libro.get("activo")
                is_activo = (
                    activo_value == 1
                    or activo_value is True
                    or str(activo_value).strip().lower() in ("1", "true")
                )
                cat_val = libro.get("id_categoria") if libro.get("id_categoria") is not None else libro.get("categoria")
                cat_name = categorias_map.get(str(cat_val), libro.get("categoria") or "")
                ws.append([
                    (libro.get("codigo_barras") or ""),
                    (libro.get("titulo") or ""),
                    (libro.get("clasificacion_dui") or ""),
                    (cat_name),
                    (libro.get("autores") or ""),
                    (libro.get("isbn") or ""),
                    (libro.get("tipo") or ""),
                    ("ACTIVO" if is_activo else "INACTIVO"),
                ])

            nombre = f"libros_{estado.lower()}_{nombre_sufijo}.xlsx"
            out_path = os.path.join(tempfile.gettempdir(), nombre)
            try:
                wb.save(out_path)
                snack(f"Excel generado: {nombre}")
                open_file_default_app(out_path)
            except Exception as ex:
                snack(f"No se pudo generar Excel: {ex}")

        def export_semana_actual(estado: str):
            i, f_ = rango_semana_actual()
            exportar_excel_estado_rango(estado, i, f_, "semana_actual")

        def export_semana_anterior(estado: str):
            i, f_ = rango_semana_anterior()
            exportar_excel_estado_rango(estado, i, f_, "semana_anterior")

        def export_mes_actual(estado: str):
            i, f_ = rango_mes_actual()
            exportar_excel_estado_rango(estado, i, f_, "mes_actual")

        def export_mes_anterior(estado: str):
            i, f_ = rango_mes_anterior()
            exportar_excel_estado_rango(estado, i, f_, "mes_anterior")

        def abrir_dialogo_mes(estado: str):
            ahora = datetime.now()
            meses = [
                (1, "Enero"), (2, "Febrero"), (3, "Marzo"), (4, "Abril"),
                (5, "Mayo"), (6, "Junio"), (7, "Julio"), (8, "Agosto"),
                (9, "Septiembre"), (10, "Octubre"), (11, "Noviembre"), (12, "Diciembre"),
            ]

            mes_dd = ft.Dropdown(
                label="Mes",
                width=180,
                options=[ft.dropdown.Option(key=str(m), text=nombre) for m, nombre in meses],
                value=str(ahora.month),
            )

            anios = [ahora.year - 3 + i for i in range(7)]
            anio_dd = ft.Dropdown(
                label="Año",
                width=140,
                options=[ft.dropdown.Option(key=str(y), text=str(y)) for y in anios],
                value=str(ahora.year),
            )

            def confirmar(e):
                try:
                    y = int(anio_dd.value)
                    m = int(mes_dd.value)
                    i, f_ = rango_mes_especifico(y, m)
                    suf = f"{y}_{m:02d}"
                    exportar_excel_estado_rango(estado, i, f_, f"mes_{suf}")
                    dlg.open = False
                    self._page.update()
                except Exception as ex:
                    snack(f"Error al exportar: {ex}")

            dlg = ft.AlertDialog(
                modal=True,
                content=ft.Container(
                    width=400,
                    padding=10,
                    content=ft.Column([
                        ft.Text("Elegir mes para exportar", weight=ft.FontWeight.BOLD),
                        ft.Row([mes_dd, anio_dd], spacing=8),
                        ft.Row([
                            ft.TextButton("Cancelar", on_click=lambda e: setattr(dlg, 'open', False) or self._page.update()),
                            ft.ElevatedButton("Exportar", on_click=confirmar),
                        ], alignment=ft.MainAxisAlignment.END)
                    ], spacing=10)
                )
            )

            self._page.overlay.append(dlg)
            dlg.open = True
            self._page.update()

        # Menús de exportación eliminados; queda solo el botón principal `Descargar Excel`.

        # Botón para descargar TODO el Excel (incluye registros no visibles en la tabla)
        def exportar_excel_todos(e=None):
            if not self._libros_cache:
                self._libros_cache = self._db.get_libros()

            libros = list(self._libros_cache or [])

            # Recolectar todas las claves disponibles en los registros
            all_keys = set()
            for l in libros:
                if isinstance(l, dict):
                    all_keys.update(l.keys())

            # Quitar campos de descripción asociados a categoría (p. ej. 'descripcion_categoria')
            for _k in list(all_keys):
                lk = _k.lower()
                if "categoria" in lk and "descripcion" in lk:
                    all_keys.discard(_k)

            # Orden preferido (si existen)
            preferred = [
                "codigo_barras", "titulo", "clasificacion_dui", "categoria",
                "autores", "isbn", "tipo", "activo",
            ]

            # Excluir el identificador interno si existe
            if "id_libro" in all_keys:
                all_keys.discard("id_libro")

            keys = [k for k in preferred if k in all_keys]
            # Añadir el resto de claves en orden alfabético para estabilidad
            rest = sorted([k for k in all_keys if k not in keys])
            keys.extend(rest)

            # Construir encabezados legibles
            def human(k: str) -> str:
                text = k.replace("_", " ").title()
                # Cambiar anio/año de publicacion a Año De Publicación
                if "anio" in text.lower() and "publicacion" in text.lower():
                    text = "Año De Publicación"
                return text

            headers = [human(k) for k in keys]
            # Asegurar que la columna de categoría use exactamente 'Categoria'
            for idx, k in enumerate(keys):
                if str(k).lower() in ("categoria", "id_categoria", "idcategoria"):
                    headers[idx] = "Categoria"

            wb = Workbook()
            ws = wb.active
            ws.title = "Todos"
            
            # Importar estilos
            from openpyxl.styles import Border, Side, PatternFill, Font, Alignment
            
            # Estilos
            header_fill = PatternFill(start_color="1976d2", end_color="1976d2", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # Bordes
            thin_border = Border(
                left=Side(style='thin', color="cccccc"),
                right=Side(style='thin', color="cccccc"),
                top=Side(style='thin', color="cccccc"),
                bottom=Side(style='thin', color="cccccc")
            )
            
            # Rellenos alternados
            row_fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            row_fill_light = PatternFill(start_color="f5f7fa", end_color="f5f7fa", fill_type="solid")
            
            # Agregar encabezados con estilo
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = thin_border
            
            ws.row_dimensions[1].height = 25

            import json

            # Preparar mapeo id_categoria -> nombre (si aplica)
            categorias_map = {}
            try:
                cats = self._db.get_categorias()
                for c in cats:
                    # pueden venir como id_categoria o id
                    cid = c.get("id_categoria") or c.get("id")
                    if cid is not None:
                        categorias_map[str(cid)] = c.get("nombre") or c.get("nombre_categoria") or c.get("nombre_categoria", "")
            except Exception:
                categorias_map = {}

            def serialize(v, key=None):
                if v is None:
                    return ""
                # Mapear categoría por id si es necesario
                if key and key.lower() in ("id_categoria", "idcategoria", "categoria"):
                    try:
                        # si viene un entero o texto numérico
                        sk = str(v)
                        if sk in categorias_map:
                            return categorias_map[sk]
                    except Exception:
                        pass

                if isinstance(v, (list, tuple)):
                    return ", ".join([str(x) for x in v])
                if isinstance(v, dict):
                    try:
                        return json.dumps(v, ensure_ascii=False)
                    except Exception:
                        return str(v)
                if isinstance(v, bytes):
                    return "<BINARY>"
                if isinstance(v, datetime):
                    return v.isoformat()
                return str(v)

            for row_idx, libro in enumerate(libros, 2):
                row_data = [serialize(libro.get(k), k) if isinstance(libro, dict) else "" for k in keys]
                
                # Determinar relleno (alternado)
                current_fill = row_fill_light if row_idx % 2 == 0 else row_fill_white
                
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = value
                    cell.fill = current_fill
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            # Ajustar ancho de columnas automáticamente
            for col_idx, header in enumerate(headers, 1):
                ws.column_dimensions[chr(64 + col_idx)].width = min(40, max(12, len(str(header)) + 2))

            nombre = f"libros_todos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            out_path = os.path.join(tempfile.gettempdir(), nombre)
            try:
                wb.save(out_path)
                snack(f"Excel generado: {nombre}")
                open_file_default_app(out_path)
            except Exception as ex:
                snack(f"No se pudo generar Excel: {ex}")

        btn_descargar_excel = ft.ElevatedButton(
            content=ft.Row([ft.Icon(ft.Icons.DOWNLOAD), ft.Text("Descargar Excel")], spacing=8),
            on_click=exportar_excel_todos,
            bgcolor="#1976d2",
            color=ft.Colors.WHITE,
            height=48,
            style=ft.ButtonStyle(
                shape=ft.RoundedRectangleBorder(radius=10),
            ),
        )

        btn_refrescar = ft.IconButton(
            icon=ft.Icons.REFRESH_ROUNDED,
            icon_color=ft.Colors.WHITE,
            bgcolor="#1976d2",
            tooltip="Refrescar datos",
            on_click=lambda e: self._mostrar_libros() if hasattr(self, '_mostrar_libros') else None,
            icon_size=24,
            height=48,
            width=48,
            style=ft.ButtonStyle(
                shape=ft.RoundedRectangleBorder(radius=10),
            ),
        )

        header = ft.Container(
            content=ft.Row(
                [
                    ft.Row([
                        ft.Icon(ft.Icons.ASSIGNMENT_ROUNDED, color="#1565c0", size=32),
                        ft.Text("Libros Registrados", size=26, weight=ft.FontWeight.BOLD, color="#263238"),
                    ], spacing=12),
                    ft.Row([
                        btn_refrescar,
                        button_crear_libro,
                    ], spacing=12),
                ],
                alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
            ),
            padding=ft.padding.symmetric(horizontal=30, vertical=20),
            margin=ft.margin.symmetric(horizontal=30),
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            shadow=ft.BoxShadow(
                spread_radius=0,
                blur_radius=8,
                color=ft.Colors.with_opacity(0.08, ft.Colors.BLACK),
                offset=ft.Offset(0, 2),
            ),
        )

        # --- Métricas (estilo copiado de reservas)
        self.total_label = ft.Text("0", size=32, weight=ft.FontWeight.BOLD, color="#263238")
        self.activos_label = ft.Text("0", size=32, weight=ft.FontWeight.BOLD, color="#263238")
        self.digitales_label = ft.Text("0", size=32, weight=ft.FontWeight.BOLD, color="#263238")
        self.fisicos_label = ft.Text("0", size=32, weight=ft.FontWeight.BOLD, color="#263238")

        def crear_stat_card(titulo, valor_widget, icon, color, bgcolor):
            return ft.Container(
                content=ft.Column([
                    ft.Row([
                        ft.Container(
                            content=ft.Icon(icon, color=ft.Colors.WHITE, size=28),
                            bgcolor=color,
                            width=56,
                            height=56,
                            border_radius=12,
                            alignment=ft.Alignment(0, 0),
                        ),
                        ft.Column([
                            valor_widget,
                            ft.Text(titulo, size=13, color="#757575", weight=ft.FontWeight.W_500),
                        ], spacing=0, alignment=ft.MainAxisAlignment.CENTER),
                    ], alignment=ft.MainAxisAlignment.START, spacing=15),
                ], spacing=0),
                bgcolor=bgcolor,
                border=ft.border.all(1, "#e0e0e0"),
                border_radius=12,
                padding=20,
                width=250,
                shadow=ft.BoxShadow(
                    spread_radius=0,
                    blur_radius=10,
                    color=ft.Colors.with_opacity(0.1, ft.Colors.BLACK),
                    offset=ft.Offset(0, 2),
                ),
            )

        stats_row = ft.Row([
            crear_stat_card("Total de libros", self.total_label, ft.Icons.LIBRARY_BOOKS_ROUNDED, "#5e35b1", ft.Colors.WHITE),
            crear_stat_card("Activos", self.activos_label, ft.Icons.BOOK_ROUNDED, "#1976d2", ft.Colors.WHITE),
            crear_stat_card("Digitales", self.digitales_label, ft.Icons.PICTURE_AS_PDF_ROUNDED, "#00838f", ft.Colors.WHITE),
            crear_stat_card("Físicos", self.fisicos_label, ft.Icons.MENU_BOOK_ROUNDED, "#6a1b9a", ft.Colors.WHITE),
        ], spacing=20, scroll=ft.ScrollMode.AUTO)

        def actualizar_stats():
            try:
                total = len(self._libros_cache) if hasattr(self, '_libros_cache') else 0
                activos = 0
                digitales = 0
                fisicos = 0
                for l in (self._libros_cache or []):
                    av = l.get('activo')
                    is_activo = (av == 1 or av is True or str(av).strip().lower() in ('1', 'true'))
                    if is_activo:
                        activos += 1
                    tipo = (l.get('tipo') or '').strip().upper()
                    if tipo == 'DIGITAL':
                        digitales += 1
                    elif tipo == 'FISICO':
                        fisicos += 1
                self.total_label.value = str(total)
                self.activos_label.value = str(activos)
                self.digitales_label.value = str(digitales)
                self.fisicos_label.value = str(fisicos)
                if hasattr(self, '_page'):
                    self._page.update()
            except Exception:
                pass

        # Contenedor principal con header, métricas, filtros y tabla
        filtros_bar = ft.Container(
            content=ft.Row(
                [
                    search_input,
                    estado_dd,
                    tipo_dd,
                    ft.Row([btn_descargar_excel], spacing=8),
                ],
                alignment=ft.MainAxisAlignment.START,
                spacing=15,
            ),
            padding=ft.padding.symmetric(horizontal=30, vertical=15),
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            shadow=ft.BoxShadow(
                spread_radius=0,
                blur_radius=8,
                color=ft.Colors.with_opacity(0.08, ft.Colors.BLACK),
                offset=ft.Offset(0, 2),
            ),
        )

        tabla_container = ft.Container(
            content=ft.Column([
                ft.Container(
                    content=libros_table,
                    alignment=ft.Alignment.CENTER,
                ),
                ft.Container(
                    content=ft.Row(
                        [btn_prev, pagination_label, btn_next],
                        alignment=ft.MainAxisAlignment.CENTER,
                        spacing=6,
                    ),
                    padding=ft.padding.only(top=8),
                ),
            ], scroll=ft.ScrollMode.AUTO),
            padding=20,
            bgcolor=ft.Colors.WHITE,
            border_radius=12,
            shadow=ft.BoxShadow(
                spread_radius=0,
                blur_radius=8,
                color=ft.Colors.with_opacity(0.08, ft.Colors.BLACK),
                offset=ft.Offset(0, 2),
            ),
            margin=ft.margin.symmetric(horizontal=30),
            expand=True,
        )

        # Layout principal (match reservas)
        self.controls = [
            ft.Container(
                content=ft.Column(
                    [
                        header,
                        ft.Container(
                            content=stats_row,
                            padding=ft.padding.symmetric(horizontal=30),
                        ),
                        ft.Container(
                            content=filtros_bar,
                            padding=ft.padding.symmetric(horizontal=30),
                        ),
                        ft.Container(
                            content=tabla_container,
                            padding=0,
                            alignment=ft.Alignment.CENTER,
                            expand=True,
                        ),
                    ],
                    spacing=20,
                    expand=True,
                ),
                bgcolor="#f5f7fa",
                padding=ft.padding.symmetric(vertical=20),
                expand=True,
            )
        ]

        mostrar_libros()

    # ===========================
    # Métodos auxiliares para autores (Crear)
    # ===========================
    def actualizar_dropdown_autores_crear(self):
        self.autores_dropdown.options = [
            ft.dropdown.Option(key=str(a.get("id_autor")), text=a.get("nombre_completo", "Sin nombre"))
            for a in self.autores_todos
        ]
        if hasattr(self, '_page'):
            self._page.update()
    
    def filtrar_autores_crear(self, e):
        busqueda = (self.buscar_autor_input.value or "").strip().lower()
        if not busqueda:
            self.actualizar_dropdown_autores_crear()
            return
        
        autores_filtrados = [
            a for a in self.autores_todos
            if busqueda in a.get("nombre_completo", "").lower()
        ]
        
        self.autores_dropdown.options = [
            ft.dropdown.Option(key=str(a.get("id_autor")), text=a.get("nombre_completo", "Sin nombre"))
            for a in autores_filtrados
        ]
        self._page.update()
    
    def agregar_autor_crear(self, e):
        if not self.autores_dropdown.value:
            return
        
        autor_id = self.autores_dropdown.value
        # Verificar si ya está agregado
        if any(a["id"] == autor_id for a in self.autores_seleccionados):
            return
        
        # Buscar el nombre del autor
        autor = next((a for a in self.autores_todos if str(a.get("id_autor")) == autor_id), None)
        if not autor:
            return
        
        autor_nombre = autor.get("nombre_completo", "Sin nombre")
        self.autores_seleccionados.append({"id": autor_id, "nombre": autor_nombre})
        
        # Actualizar la vista
        self.actualizar_lista_autores_seleccionados_crear()
    
    def remover_autor_crear(self, autor_id):
        self.autores_seleccionados = [a for a in self.autores_seleccionados if a["id"] != autor_id]
        self.actualizar_lista_autores_seleccionados_crear()
    
    def actualizar_lista_autores_seleccionados_crear(self):
        self.autores_seleccionados_column.controls.clear()
        
        for autor in self.autores_seleccionados:
            self.autores_seleccionados_column.controls.append(
                ft.Container(
                    content=ft.Row([
                        ft.Text(autor["nombre"], size=13),
                        ft.IconButton(
                            icon=ft.Icons.CLOSE,
                            icon_size=16,
                            icon_color=ft.Colors.RED,
                            tooltip="Remover",
                            on_click=lambda e, aid=autor["id"]: self.remover_autor_crear(aid),
                        ),
                    ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    padding=8,
                    bgcolor="#e3f2fd",
                    border_radius=6,
                )
            )
        
        self._page.update()
    
    # Helper methods para editar libro - Sistema de autores
    def actualizar_dropdown_autores_editar(self):
        """Actualiza el dropdown con la lista completa de autores para editar"""
        self.autores_dropdown.options = [
            ft.dropdown.Option(key=str(autor.get("id_autor")), text=autor.get("nombre_completo", "Sin nombre"))
            for autor in self.autores_todos
        ]
    
    def filtrar_autores_editar(self, e):
        """Filtra el dropdown de autores según el texto de búsqueda para editar"""
        busqueda = self.buscar_autor_input.value.lower()
        if busqueda:
            self.autores_dropdown.options = [
                ft.dropdown.Option(key=str(autor.get("id_autor")), text=autor.get("nombre_completo", "Sin nombre"))
                for autor in self.autores_todos
                if busqueda in autor.get("nombre_completo", "").lower()
            ]
        else:
            self.actualizar_dropdown_autores_editar()
        
        self._page.update()
    
    def agregar_autor_editar(self, e):
        """Agrega el autor seleccionado en el dropdown a la lista de seleccionados para editar"""
        if not self.autores_dropdown.value:
            return
        
        autor_id = self.autores_dropdown.value
        
        # Verificar si ya está agregado
        if any(a["id"] == autor_id for a in self.autores_seleccionados):
            self.snack("Este autor ya fue agregado", "warning")
            return
        
        # Buscar el nombre del autor
        autor = next((a for a in self.autores_todos if str(a.get("id_autor")) == autor_id), None)
        if autor:
            self.autores_seleccionados.append({
                "id": autor_id,
                "nombre": autor.get("nombre_completo", "Sin nombre")
            })
            self.actualizar_lista_autores_seleccionados_editar()
            self.buscar_autor_input.value = ""
            self.autores_dropdown.value = None
            self.actualizar_dropdown_autores_editar()
    
    def remover_autor_editar(self, autor_id):
        """Remueve un autor de la lista de seleccionados para editar"""
        self.autores_seleccionados = [a for a in self.autores_seleccionados if a["id"] != autor_id]
        self.actualizar_lista_autores_seleccionados_editar()
    
    def actualizar_lista_autores_seleccionados_editar(self):
        """Actualiza la lista visual de autores seleccionados para editar"""
        self.autores_seleccionados_column.controls.clear()
        
        for autor in self.autores_seleccionados:
            self.autores_seleccionados_column.controls.append(
                ft.Container(
                    content=ft.Row([
                        ft.Text(autor["nombre"], size=13),
                        ft.IconButton(
                            icon=ft.Icons.CLOSE,
                            icon_size=16,
                            icon_color=ft.Colors.RED,
                            tooltip="Remover",
                            on_click=lambda e, aid=autor["id"]: self.remover_autor_editar(aid),
                        ),
                    ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    padding=8,
                    bgcolor="#e3f2fd",
                    border_radius=6,
                )
            )
        
        self._page.update()

    # ===========================
    # Crear Libro con Diálogo
    # ===========================
    def abrir_dialogo_crear_libro(self, e):
        INPUT_STYLE = dict(
            width=320,
            height=44,
            bgcolor="#f5f7fa",
            border_radius=8,
            border_color="#cfd8dc",
            focused_border_color="#1976d2",
            text_size=14,
        )

        # Campos del formulario
        self.titulo_input = ft.TextField(label="Título *", autofocus=True, **INPUT_STYLE)
        self.isbn_input = ft.TextField(label="ISBN", **INPUT_STYLE)
        self.codigo_barras_input = ft.TextField(label="Código de barras", **INPUT_STYLE)
        self.anio_input = ft.TextField(label="Año de publicación", keyboard_type=ft.KeyboardType.NUMBER, **INPUT_STYLE)
        self.edicion_input = ft.TextField(label="Edición", **INPUT_STYLE)
        
        self.tipo_dd = ft.Dropdown(
            label="Tipo *",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            options=[
                ft.dropdown.Option("FISICO"),
                ft.dropdown.Option("DIGITAL"),
            ],
        )

        self.descripcion_input = ft.TextField(
            label="Descripción",
            multiline=True,
            min_lines=5,
            max_lines=8,
            width=680,
            bgcolor="#f5f7fa",
            border_radius=8,
        )

        cats = self._db.get_categorias()
        self.categoria_dd = ft.Dropdown(
            label="Categoría *",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            options=[
                ft.dropdown.Option(key=str(c["id_categoria"]), text=c["nombre"])
                for c in cats
            ],
        )

        self.clasificacion_dui_input = ft.TextField(label="Clasificación DUI *", **INPUT_STYLE)
        
        # Sistema de búsqueda y selección de autores
        self.autores_todos = self._db.get_autores()
        self.autores_seleccionados = []  # Lista de {id, nombre}
        
        self.buscar_autor_input = ft.TextField(
            label="Buscar autor",
            width=320,
            height=44,
            bgcolor="#f5f7fa",
            border_radius=8,
            border_color="#cfd8dc",
            text_size=14,
            on_change=self.filtrar_autores_crear,
        )
        
        self.autores_dropdown = ft.Dropdown(
            label="Selecciona un autor",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            options=[],
        )
        
        self.btn_agregar_autor = ft.IconButton(
            icon=ft.Icons.ADD,
            bgcolor="#1976d2",
            icon_color=ft.Colors.WHITE,
            tooltip="Agregar autor",
            on_click=self.agregar_autor_crear,
        )
        
        self.autores_seleccionados_column = ft.Column(spacing=4)
        
        self.autores_container = ft.Container(
            content=ft.Column([
                ft.Row([self.buscar_autor_input, self.autores_dropdown, self.btn_agregar_autor], spacing=8),
                ft.Divider(height=8),
                ft.Text("Autores seleccionados:", size=12, weight=ft.FontWeight.BOLD, color="#666"),
                self.autores_seleccionados_column,
            ], spacing=8),
            padding=10,
            border=ft.border.all(1, "#cfd8dc"),
            border_radius=8,
            bgcolor="#fafafa",
        )
        
        # Inicializar dropdown con todos los autores
        self.actualizar_dropdown_autores_crear()

        self.btn_guardar = ft.ElevatedButton(
            content=ft.Row([ft.Icon(ft.Icons.CHECK), ft.Text("Guardar")], spacing=8),
            bgcolor="#1976d2",
            color=ft.Colors.WHITE,
            on_click=self.crear_libro,
        )

        header_row = ft.Row([
            ft.Row([
                ft.Container(
                    content=ft.Icon(ft.Icons.BOOK, size=22, color="#1B6F7A"),
                    bgcolor=ft.Colors.with_opacity(0.06, ft.Colors.GREEN),
                    width=40,
                    height=40,
                    border_radius=8,
                    alignment=ft.Alignment.CENTER,
                ),
                ft.Column([
                    ft.Text("Crear libro", size=18, weight=ft.FontWeight.BOLD),
                    ft.Text("Completa los datos del libro", size=12, color="#666")
                ], spacing=2)
            ], spacing=10),
            ft.Container(
                content=ft.Icon(ft.Icons.CLOSE, size=18, color="#666"),
                width=36,
                height=36,
                alignment=ft.Alignment.CENTER,
                on_click=self.cerrar_dialogo_libro,
                border_radius=8
            ),
        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN)

        form_column = ft.Column([
            # Sección: Información Básica
            ft.Text("Información Básica", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            self.titulo_input,
            ft.Row([self.tipo_dd, self.categoria_dd], spacing=12),
            
            # Sección: Identificadores
            ft.Text("Identificadores", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            self.clasificacion_dui_input,
            ft.Row([self.isbn_input, self.codigo_barras_input], spacing=12),
            
            # Sección: Detalles de Publicación
            ft.Text("Detalles de Publicación", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            ft.Row([self.anio_input, self.edicion_input], spacing=12),
            
            # Sección: Descripción y Autores
            ft.Text("Descripción y Autores", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            ft.Text("Selecciona los autores:", size=12, color="#666"),
            self.autores_container,
            self.descripcion_input,
            
        ], spacing=12, horizontal_alignment=ft.CrossAxisAlignment.CENTER)

        content = ft.Column([
            header_row,
            ft.Divider(height=8, color="transparent"),
            form_column,
            ft.Divider(height=6, color="transparent"),
            ft.Row([ft.ElevatedButton("Cancelar", on_click=self.cerrar_dialogo_libro, bgcolor="#757575", color=ft.Colors.WHITE), self.btn_guardar], alignment=ft.MainAxisAlignment.END, spacing=12)
        ], spacing=10, scroll=ft.ScrollMode.AUTO)

        self.dialog = ft.AlertDialog(
            modal=True,
            content=ft.Container(width=780, padding=ft.padding.all(18), bgcolor=ft.Colors.WHITE, border_radius=12, content=content)
        )

        self._page.overlay.clear()
        self._page.overlay.append(self.dialog)
        self.dialog.open = True
        self._page.update()

    # ===========================
    # Editar Libro con Diálogo
    # ===========================
    def abrir_dialogo_editar_libro(self, id_libro: int):
        libro = self._db.get_libro_detalle(int(id_libro))
        if not libro:
            self._page.snack_bar = ft.SnackBar(ft.Text("Libro no encontrado"), bgcolor="#b71c1c")
            self._page.snack_bar.open = True
            self._page.update()
            return

        self._edit_libro_id = int(id_libro)

        INPUT_STYLE = dict(
            width=320,
            height=44,
            bgcolor="#f5f7fa",
            border_radius=8,
            border_color="#cfd8dc",
            focused_border_color="#1976d2",
            text_size=14,
        )

        self.titulo_input = ft.TextField(label="Título *", value=libro.get("titulo", ""), **INPUT_STYLE)
        self.isbn_input = ft.TextField(label="ISBN", value=libro.get("isbn", "") or "", **INPUT_STYLE)
        self.codigo_barras_input = ft.TextField(
            label="Código de barras",
            value=libro.get("codigo_barras", "") or "",
            **INPUT_STYLE,
        )
        self.anio_input = ft.TextField(
            label="Año de publicación",
            value=str(libro.get("anio_publicacion", "")) if libro.get("anio_publicacion") else "",
            keyboard_type=ft.KeyboardType.NUMBER,
            **INPUT_STYLE,
        )
        self.edicion_input = ft.TextField(label="Edición", value=libro.get("edicion", "") or "", **INPUT_STYLE)

        self.tipo_dd = ft.Dropdown(
            label="Tipo *",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            value=(libro.get("tipo", "") or "").upper(),
            options=[
                ft.dropdown.Option("FISICO"),
                ft.dropdown.Option("DIGITAL"),
            ],
        )

        self.descripcion_input = ft.TextField(
            label="Descripción",
            value=libro.get("descripcion", "") or "",
            multiline=True,
            min_lines=5,
            max_lines=8,
            width=680,
            bgcolor="#f5f7fa",
            border_radius=8,
        )

        cats = self._db.get_categorias()
        self.categoria_dd = ft.Dropdown(
            label="Categoría *",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            value=str(libro.get("id_categoria", "")),
            options=[
                ft.dropdown.Option(key=str(c["id_categoria"]), text=c["nombre"])
                for c in cats
            ],
        )

        self.clasificacion_dui_input = ft.TextField(
            label="Clasificación DUI *",
            value=libro.get("clasificacion_dui", "") or "",
            **INPUT_STYLE,
        )
        
        # Sistema de búsqueda y selección de autores para edición
        self.autores_todos = self._db.get_autores()
        
        # Obtener los IDs de los autores del libro desde la base de datos
        autores_ids_libro = self._db.get_autores_libro(int(id_libro))
        autores_ids_list = [id.strip() for id in autores_ids_libro.split(",") if id.strip()]
        
        # Inicializar autores seleccionados con los del libro
        self.autores_seleccionados = []
        for autor_id in autores_ids_list:
            autor = next((a for a in self.autores_todos if str(a.get("id_autor")) == autor_id), None)
            if autor:
                self.autores_seleccionados.append({
                    "id": str(autor.get("id_autor")),
                    "nombre": autor.get("nombre_completo", "Sin nombre")
                })
        
        self.buscar_autor_input = ft.TextField(
            label="Buscar autor",
            width=320,
            height=44,
            bgcolor="#f5f7fa",
            border_radius=8,
            border_color="#cfd8dc",
            text_size=14,
            on_change=self.filtrar_autores_editar,
        )
        
        self.autores_dropdown = ft.Dropdown(
            label="Selecciona un autor",
            width=320,
            bgcolor="#f5f7fa",
            border_radius=8,
            options=[],
        )
        
        self.btn_agregar_autor = ft.IconButton(
            icon=ft.Icons.ADD,
            bgcolor="#1976d2",
            icon_color=ft.Colors.WHITE,
            tooltip="Agregar autor",
            on_click=self.agregar_autor_editar,
        )
        
        self.autores_seleccionados_column = ft.Column(spacing=4)
        
        self.autores_container = ft.Container(
            content=ft.Column([
                ft.Row([self.buscar_autor_input, self.autores_dropdown, self.btn_agregar_autor], spacing=8),
                ft.Divider(height=8),
                ft.Text("Autores seleccionados:", size=12, weight=ft.FontWeight.BOLD, color="#666"),
                self.autores_seleccionados_column,
            ], spacing=8),
            padding=10,
            border=ft.border.all(1, "#cfd8dc"),
            border_radius=8,
            bgcolor="#fafafa",
        )
        
        # Inicializar dropdown y lista de seleccionados
        self.actualizar_dropdown_autores_editar()
        self.actualizar_lista_autores_seleccionados_editar()

        self.btn_guardar = ft.ElevatedButton(
            content=ft.Row([ft.Icon(ft.Icons.CHECK), ft.Text("Guardar")], spacing=8),
            bgcolor="#1976d2",
            color=ft.Colors.WHITE,
            on_click=self.guardar_edicion_libro,
        )

        header_row = ft.Row([
            ft.Row([
                ft.Container(
                    content=ft.Icon(ft.Icons.EDIT, size=22, color="#1B6F7A"),
                    bgcolor=ft.Colors.with_opacity(0.06, ft.Colors.GREEN),
                    width=40,
                    height=40,
                    border_radius=8,
                    alignment=ft.Alignment.CENTER,
                ),
                ft.Column([
                    ft.Text("Editar libro", size=18, weight=ft.FontWeight.BOLD),
                    ft.Text("Actualiza los datos del libro", size=12, color="#666")
                ], spacing=2)
            ], spacing=10),
            ft.Container(
                content=ft.Icon(ft.Icons.CLOSE, size=18, color="#666"),
                width=36,
                height=36,
                alignment=ft.Alignment.CENTER,
                on_click=self.cerrar_dialogo_libro,
                border_radius=8,
            ),
        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN)

        form_column = ft.Column([
            ft.Text("Información Básica", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            self.titulo_input,
            ft.Row([self.tipo_dd, self.categoria_dd], spacing=12),

            ft.Text("Identificadores", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            self.clasificacion_dui_input,
            ft.Row([self.isbn_input, self.codigo_barras_input], spacing=12),

            ft.Text("Detalles de Publicación", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            ft.Row([self.anio_input, self.edicion_input], spacing=12),

            ft.Text("Descripción y Autores", size=14, weight=ft.FontWeight.BOLD, color="#0b495c"),
            ft.Text("Selecciona los autores:", size=12, color="#666"),
            self.autores_container,
            self.descripcion_input,
        ], spacing=12, horizontal_alignment=ft.CrossAxisAlignment.CENTER)

        content = ft.Column([
            header_row,
            ft.Divider(height=8, color="transparent"),
            form_column,
            ft.Divider(height=6, color="transparent"),
            ft.Row([
                ft.ElevatedButton("Cancelar", on_click=self.cerrar_dialogo_libro, bgcolor="#757575", color=ft.Colors.WHITE),
                self.btn_guardar
            ], alignment=ft.MainAxisAlignment.END, spacing=12)
        ], spacing=10, scroll=ft.ScrollMode.AUTO)

        self.dialog = ft.AlertDialog(
            modal=True,
            content=ft.Container(
                width=780,
                padding=ft.padding.all(18),
                bgcolor=ft.Colors.WHITE,
                border_radius=12,
                content=content,
            )
        )

        self._page.overlay.clear()
        self._page.overlay.append(self.dialog)
        self.dialog.open = True
        self._page.update()

    def guardar_edicion_libro(self, e):
        try:
            titulo = (self.titulo_input.value or "").strip()
            tipo = (self.tipo_dd.value or "").strip()

            if not titulo:
                self._page.snack_bar = ft.SnackBar(ft.Text("El título es obligatorio"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            if not tipo:
                self._page.snack_bar = ft.SnackBar(ft.Text("El tipo es obligatorio"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            if not self.categoria_dd.value:
                self._page.snack_bar = ft.SnackBar(ft.Text("Debe seleccionar una categoría"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            clasificacion_dui = (self.clasificacion_dui_input.value or "").strip()
            if not clasificacion_dui:
                self._page.snack_bar = ft.SnackBar(ft.Text("La Clasificación DUI es obligatoria"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            anio = (self.anio_input.value or "").strip()
            anio_publicacion = int(anio) if anio else None

            # Obtener IDs de autores seleccionados del nuevo sistema
            autores_csv = ",".join([a["id"] for a in self.autores_seleccionados])
            codigo_barras = (self.codigo_barras_input.value or "").strip()

            self._db.update_libro(
                id_libro=self._edit_libro_id,
                titulo=titulo,
                isbn=(self.isbn_input.value or "").strip(),
                anio_publicacion=anio_publicacion,
                edicion=(self.edicion_input.value or "").strip(),
                tipo=tipo,
                descripcion=(self.descripcion_input.value or "").strip(),
                id_categoria=int(self.categoria_dd.value),
                autores_ids_csv=autores_csv if autores_csv else None,
                clasificacion_dui=clasificacion_dui,
                codigo_barras=codigo_barras if codigo_barras else None,
            )

            self._page.snack_bar = ft.SnackBar(ft.Text("✅ Cambios guardados correctamente"), bgcolor="#1b5e20")
            self._page.snack_bar.open = True
            self.cerrar_dialogo_libro()
            if hasattr(self, "_mostrar_libros_sin_reset"):
                self._mostrar_libros_sin_reset()

        except Exception as ex:
            self._page.snack_bar = ft.SnackBar(ft.Text(f"Error: {str(ex)[:150]}"), bgcolor="#b71c1c")
            self._page.snack_bar.open = True
            self._page.update()

    def crear_libro(self, e):
        try:
            print("[CREATE] Iniciando crear_libro...")
            titulo = (self.titulo_input.value or "").strip()
            tipo = (self.tipo_dd.value or "").strip()

            if not titulo:
                self._page.snack_bar = ft.SnackBar(ft.Text("El título es obligatorio"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            if not tipo:
                self._page.snack_bar = ft.SnackBar(ft.Text("El tipo es obligatorio"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            if not self.categoria_dd.value:
                self._page.snack_bar = ft.SnackBar(ft.Text("Debe seleccionar una categoría"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            clasificacion_dui = (self.clasificacion_dui_input.value or "").strip()
            if not clasificacion_dui:
                self._page.snack_bar = ft.SnackBar(ft.Text("La Clasificación DUI es obligatoria"), bgcolor="#b71c1c")
                self._page.snack_bar.open = True
                self._page.update()
                return

            anio = (self.anio_input.value or "").strip()
            anio_publicacion = int(anio) if anio else None

            # Obtener IDs de autores seleccionados
            autores_csv = ",".join([a["id"] for a in self.autores_seleccionados])

            codigo_barras = (self.codigo_barras_input.value or "").strip()

            print(f"[CREATE] Parámetros:")
            print(f"  - titulo: {titulo}")
            print(f"  - tipo: {tipo}")
            print(f"  - id_categoria: {self.categoria_dd.value}")
            print(f"  - clasificacion_dui: {clasificacion_dui}")
            print(f"  - anio_publicacion: {anio_publicacion}")

            print("[CREATE] Llamando a _db.crear_libro()...")
            resultado = self._db.crear_libro(
                titulo=titulo,
                isbn=(self.isbn_input.value or "").strip(),
                anio_publicacion=anio_publicacion,
                edicion=(self.edicion_input.value or "").strip(),
                tipo=tipo,
                descripcion=(self.descripcion_input.value or "").strip(),
                id_categoria=int(self.categoria_dd.value),
                activo=1,
                autores_ids_csv=autores_csv if autores_csv else None,
                clasificacion_dui=clasificacion_dui,
                codigo_barras=codigo_barras if codigo_barras else None,
            )
            print(f"[CREATE] Resultado: {resultado}")

            self._page.snack_bar = ft.SnackBar(ft.Text("✅ Libro guardado correctamente"), bgcolor="#1b5e20")
            self._page.snack_bar.open = True
            self.cerrar_dialogo_libro()

            if hasattr(self, "_mostrar_libros_sin_reset"):
                self._mostrar_libros_sin_reset()
            
        except Exception as ex:
            print(f"[CREATE] ERROR: {ex}")
            import traceback
            traceback.print_exc()
            self._page.snack_bar = ft.SnackBar(ft.Text(f"Error: {str(ex)[:150]}"), bgcolor="#b71c1c")
            self._page.snack_bar.open = True
            self._page.update()

    def cerrar_dialogo_libro(self, e=None):
        if hasattr(self, 'dialog') and self.dialog:
            self.dialog.open = False
            self._page.update()
